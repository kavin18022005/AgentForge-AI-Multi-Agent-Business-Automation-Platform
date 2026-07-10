import os
import re
import json
import uuid
import asyncio
from datetime import datetime
from loguru import logger

from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.schema import HumanMessage
from langgraph.graph import StateGraph, END

from app.config import settings
from app.database import AsyncSessionLocal
from app.models.project import Project
from app.models.task import Task
from app.models.report import Report
from app.models.activity import Notification, Upload, Activity
from app.models.user import User
from sqlalchemy import select

from app.agents.state import AgentForgeState
from app.agents.base import (
    get_llm,
    extract_json,
    broadcast_ws,
    update_project_status,
    create_task_record,
    complete_task_record,
)
from app.prompts import templates

# ──────────────────────────────────────────────────────────────────────
# File Text Extraction Utilities
# ──────────────────────────────────────────────────────────────────────

def extract_text_from_file(file_path: str, file_type: str) -> str:
    """Extract text content from various file formats."""
    if not os.path.exists(file_path):
        return ""
    
    try:
        if file_type == "txt" or file_type == "csv":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
                
        elif file_type == "pdf":
            import pypdf
            reader = pypdf.PdfReader(file_path)
            text_parts = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
            return "\n".join(text_parts)
            
        elif file_type == "docx":
            import docx
            doc = docx.Document(file_path)
            text_parts = [p.text for p in doc.paragraphs]
            return "\n".join(text_parts)
            
        elif file_type == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_str = ", ".join([str(cell) for cell in row if cell is not None])
                    if row_str:
                        text_parts.append(row_str)
            return "\n".join(text_parts)
            
        elif file_type in ["png", "jpg", "jpeg"]:
            return f"[Uploaded Image File: {os.path.basename(file_path)}]"
            
    except Exception as e:
        logger.error(f"Failed to extract text from {file_path}: {e}")
        return f"[Error extracting text from file {os.path.basename(file_path)}: {str(e)}]"

    return ""

# ──────────────────────────────────────────────────────────────────────
# Agent Nodes Implementation
# ──────────────────────────────────────────────────────────────────────

async def manager_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Manager Agent"
    agent_type = "manager"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name} for project {project_id}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Analyzing business objective and preparing plan...",
        "progress": 5.0
    })
    await update_project_status(project_id, "running", 5.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Analyze goal and map execution timeline", 1)

    try:
        # Context building: user goal + uploaded documents
        goal_context = state["user_goal"]
        if state["uploaded_documents"]:
            goal_context += "\n\nUploaded Business Documents:\n" + "\n---\n".join(state["uploaded_documents"])

        prompt = templates.MANAGER_PROMPT.format(user_goal=goal_context)
        llm = get_llm(temperature=0.3)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["category"] = parsed.get("category", "strategy")
        state["execution_plan"] = parsed.get("execution_plan", {})
        # Map subtasks from plan
        state["subtasks"] = [
            {"agent": "research", "task": state["execution_plan"].get("research", "Market research")},
            {"agent": "competitor", "task": state["execution_plan"].get("competitor", "Competitor analysis")},
            {"agent": "strategy", "task": state["execution_plan"].get("strategy", "Business strategy development")},
            {"agent": "financial", "task": state["execution_plan"].get("financial", "Financial projections")},
            {"agent": "marketing", "task": state["execution_plan"].get("marketing", "Marketing and SEO plan")},
            {"agent": "risk", "task": state["execution_plan"].get("risk", "Risk assessment and register")},
        ]
        
        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 10.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def research_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Research Agent"
    agent_type = "research"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Conducting market research and demographic analysis...",
        "progress": 10.0
    })
    await update_project_status(project_id, "running", 10.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Conduct market and SWOT research", 2)

    try:
        research_focus = state.get("execution_plan", {}).get("research", "General market sizing and SWOT analysis")
        goal_context = state["user_goal"]
        if state["uploaded_documents"]:
            goal_context += "\n\nUploaded Business Documents:\n" + "\n---\n".join(state["uploaded_documents"])

        prompt = templates.RESEARCH_PROMPT.format(
            user_goal=goal_context,
            category=state.get("category", "strategy"),
            research_focus=research_focus
        )
        
        llm = get_llm(temperature=0.5)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["market_research"] = parsed
        state["industry_trends"] = [t.get("trend") for t in parsed.get("industry_trends", []) if isinstance(t, dict)]
        state["customer_segments"] = parsed.get("customer_segments", [])
        state["swot_analysis"] = parsed.get("swot_analysis", {})
        state["market_size"] = parsed.get("market_size", "TBD")

        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 25.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def competitor_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Competitor Agent"
    agent_type = "competitor"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Analyzing direct competitors and identifying market gaps...",
        "progress": 25.0
    })
    await update_project_status(project_id, "running", 25.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Identify competitors and market positioning map", 3)

    try:
        competitor_focus = state.get("execution_plan", {}).get("competitor", "Direct and indirect competitors")
        market_overview = state.get("market_research", {}).get("market_overview", "A dynamic market environment")
        
        prompt = templates.COMPETITOR_PROMPT.format(
            user_goal=state["user_goal"],
            market_overview=market_overview,
            competitor_focus=competitor_focus
        )
        
        llm = get_llm(temperature=0.5)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["competitors"] = parsed.get("competitors", [])
        state["competitive_landscape"] = parsed
        state["competitive_advantages"] = parsed.get("competitive_advantages", [])

        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 40.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def strategy_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Strategy Agent"
    agent_type = "strategy"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Developing business model canvas and roadmap...",
        "progress": 40.0
    })
    await update_project_status(project_id, "running", 40.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Synthesize Business Model Canvas and growth phases", 4)

    try:
        strategy_focus = state.get("execution_plan", {}).get("strategy", "Business model scaling")
        market_summary = state.get("market_research", {}).get("market_overview", "")
        competitor_summary = state.get("competitive_landscape", {}).get("differentiation_strategy", "")

        prompt = templates.STRATEGY_PROMPT.format(
            user_goal=state["user_goal"],
            market_summary=market_summary,
            competitive_summary=competitor_summary,
            strategy_focus=strategy_focus
        )
        
        llm = get_llm(temperature=0.4)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["business_model_canvas"] = parsed.get("business_model_canvas", {})
        state["revenue_model"] = parsed.get("revenue_model", {})
        state["pricing_strategy"] = parsed.get("revenue_model", {}).get("pricing_tiers", {})
        state["growth_strategy"] = parsed.get("growth_strategy", {})
        state["product_roadmap"] = parsed.get("product_roadmap", [])

        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 55.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def financial_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Financial Agent"
    agent_type = "financial"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Generating financial projections and ROI analysis...",
        "progress": 55.0
    })
    await update_project_status(project_id, "running", 55.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Forecast startup costs, burn rate, and break-even point", 5)

    try:
        financial_focus = state.get("execution_plan", {}).get("financial", "Profitability modeling")
        bm_canvas = state.get("business_model_canvas", {})
        rev_streams = bm_canvas.get("revenue_streams", [])
        cost_structure = bm_canvas.get("cost_structure", [])
        bm_summary = f"Revenue streams: {rev_streams}. Cost structure: {cost_structure}"

        prompt = templates.FINANCIAL_PROMPT.format(
            user_goal=state["user_goal"],
            business_model_summary=bm_summary,
            financial_focus=financial_focus
        )
        
        llm = get_llm(temperature=0.3)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["budget_estimation"] = parsed.get("startup_costs", {})
        state["revenue_forecast"] = parsed.get("revenue_forecast", {})
        state["roi_analysis"] = parsed.get("roi_analysis", {})
        state["break_even_analysis"] = parsed.get("break_even_analysis", {})
        state["financial_risks"] = parsed.get("financial_risks", [])

        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 70.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def marketing_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Marketing Agent"
    agent_type = "marketing"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Designing marketing strategies and content calendar...",
        "progress": 70.0
    })
    await update_project_status(project_id, "running", 70.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Develop SEO and channel budget allocations", 6)

    try:
        target_market = state.get("execution_plan", {}).get("target_market", "Target customers")
        comp_advantages = ", ".join(state.get("competitive_advantages", ["Quality service"]))

        prompt = templates.MARKETING_PROMPT.format(
            user_goal=state["user_goal"],
            target_market=target_market,
            competitive_advantages=comp_advantages
        )
        
        llm = get_llm(temperature=0.6)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["marketing_strategy"] = parsed
        state["seo_plan"] = parsed.get("seo_strategy", {})
        state["social_media_campaigns"] = parsed.get("social_media", {})
        state["content_calendar"] = parsed.get("content_calendar", [])
        state["brand_positioning"] = parsed.get("brand_positioning", "TBD")

        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 85.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def risk_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Risk Agent"
    agent_type = "risk"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Assessing market/operational risks and mitigation strategies...",
        "progress": 85.0
    })
    await update_project_status(project_id, "running", 85.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Build risk register and emergency response plans", 7)

    try:
        biz_context = f"Category: {state.get('category')}. Target: {state.get('market_size')}"
        financial_summary = f"startup cost total: {state.get('budget_estimation', {}).get('total')}. Risks: {state.get('financial_risks')}"

        prompt = templates.RISK_PROMPT.format(
            user_goal=state["user_goal"],
            business_context=biz_context,
            financial_summary=financial_summary
        )
        
        llm = get_llm(temperature=0.4)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["risk_register"] = parsed.get("risk_register", [])
        state["risk_score"] = float(parsed.get("overall_risk_score", 5.0))
        state["mitigation_plan"] = parsed.get("mitigation_plan", {})

        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 90.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def qa_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "QA Agent"
    agent_type = "qa"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Validating reports consistency and quality parameters...",
        "progress": 90.0
    })
    await update_project_status(project_id, "running", 90.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Analyze consistency of financials, SWOT, and strategy", 8)

    try:
        prompt = templates.QA_PROMPT.format(
            user_goal=state["user_goal"],
            research_summary=str(state.get("market_research", {})),
            competitor_summary=str(state.get("competitive_landscape", {})),
            strategy_summary=str(state.get("business_model_canvas", {})),
            financial_summary=str(state.get("roi_analysis", {})),
            marketing_summary=str(state.get("marketing_strategy", {})),
            risk_summary=str(state.get("mitigation_plan", {}))
        )
        
        llm = get_llm(temperature=0.2)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        parsed = extract_json(response.content)

        state["quality_score"] = float(parsed.get("quality_score", 8.0))
        state["qa_feedback"] = parsed.get("recommendations", [])
        state["approved"] = parsed.get("approved", True)

        await complete_task_record(task_id, parsed, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 95.0,
            "data": parsed
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state


async def report_node(state: AgentForgeState) -> AgentForgeState:
    agent_name = "Report Agent"
    agent_type = "report"
    project_id = state["project_id"]
    logger.info(f"Running node: {agent_name}")

    await broadcast_ws(project_id, {
        "type": "agent_start",
        "project_id": project_id,
        "agent": agent_type,
        "message": "Generating final executive report and export files...",
        "progress": 95.0
    })
    await update_project_status(project_id, "running", 95.0, agent_type)
    task_id = await create_task_record(project_id, agent_name, agent_type, "Publish executive report and output formats", 9)

    try:
        financial_highlights = f"Break-even: {state.get('break_even_analysis', {}).get('break_even_timeline')} months. ROI: {state.get('roi_analysis', {}).get('expected_roi_year3')} expected by Year 3."
        
        top_risks_list = []
        for r in (state.get("risk_register", []) or [])[0:3]:
            if isinstance(r, dict):
                top_risks_list.append(r.get("risk", ""))
            elif isinstance(r, str):
                top_risks_list.append(r)
        top_risks = ", ".join(top_risks_list)
        
        prompt = templates.REPORT_PROMPT.format(
            user_goal=state["user_goal"],
            project_title=state["project_title"],
            quality_score=state.get("quality_score", 8.0),
            market_summary=state.get("market_research", {}).get("market_overview", ""),
            strategy_summary=state.get("growth_strategy", {}),
            financial_highlights=financial_highlights,
            top_risks=top_risks,
            marketing_summary=state.get("brand_positioning", "")
        )
        
        llm = get_llm(temperature=0.5)
        response = await llm.ainvoke([HumanMessage(content=prompt)])
        summary_text = response.content

        state["executive_summary"] = summary_text
        state["final_report"] = {
            "manager": state.get("execution_plan"),
            "research": state.get("market_research"),
            "competitor": state.get("competitive_landscape"),
            "strategy": state.get("business_model_canvas"),
            "financial": {
                "startup_costs": state.get("budget_estimation"),
                "revenue_forecast": state.get("revenue_forecast"),
                "roi_analysis": state.get("roi_analysis"),
                "break_even_analysis": state.get("break_even_analysis"),
            },
            "marketing": state.get("marketing_strategy"),
            "risk": {
                "risk_register": state.get("risk_register"),
                "mitigation_plan": state.get("mitigation_plan"),
            },
            "qa": {
                "quality_score": state.get("quality_score"),
                "feedback": state.get("qa_feedback")
            }
        }

        # Save report database record and trigger document generation
        async with AsyncSessionLocal() as db:
            # Create a report database object
            report = Report(
                project_id=uuid.UUID(project_id),
                title=f"{state['project_title']} Intelligence Report",
                executive_summary=summary_text,
                content=state["final_report"],
                report_type=state.get("category", "business_plan"),
                status="generating",
                word_count=len(summary_text.split()),
                quality_score=state.get("quality_score", 8.0),
            )
            db.add(report)
            await db.commit()
            await db.refresh(report)
            state["report_id"] = str(report.id)

            # Generate documents synchronously using our export service
            from app.services.report_service import generate_export
            try:
                report.pdf_path = await generate_export(report, "pdf", db)
                report.docx_path = await generate_export(report, "docx", db)
                report.pptx_path = await generate_export(report, "pptx", db)
                report.markdown_path = await generate_export(report, "md", db)
                report.status = "completed"
                await db.commit()
            except Exception as e:
                logger.error(f"Exporter failed: {e}")
                report.status = "failed"
                await db.commit()

        await complete_task_record(task_id, {"report_id": state["report_id"]}, success=True)
        await broadcast_ws(project_id, {
            "type": "agent_complete",
            "project_id": project_id,
            "agent": agent_type,
            "progress": 100.0,
            "data": {"report_id": state["report_id"]}
        })
        state["completed_agents"].append(agent_type)
    except Exception as e:
        logger.error(f"Error in {agent_name}: {e}")
        state["errors"].append({"agent": agent_type, "error": str(e)})
        await complete_task_record(task_id, {"error": str(e)}, success=False)
        raise e

    return state

# ──────────────────────────────────────────────────────────────────────
# Main run_workflow implementation
# ──────────────────────────────────────────────────────────────────────

async def run_workflow(project_id: str, user_id: str):
    """Background task running the full LangGraph agent workflow."""
    logger.info(f"Kicking off multi-agent workflow for project {project_id}...")
    
    # 1. Initialize State
    async with AsyncSessionLocal() as db:
        proj_result = await db.execute(select(Project).where(Project.id == uuid.UUID(project_id)))
        project = proj_result.scalar_one_or_none()
        if not project:
            logger.error(f"Project {project_id} not found in DB.")
            return

        # Start loading files and extracting text
        uploads_result = await db.execute(select(Upload).where(Upload.project_id == uuid.UUID(project_id)))
        uploads = uploads_result.scalars().all()
        
        uploaded_docs = []
        for upload in uploads:
            if upload.analysis_status == "pending":
                upload.analysis_status = "processing"
                await db.commit()
                # Run text extraction
                extracted = extract_text_from_file(upload.file_path, upload.file_type)
                upload.extracted_text = extracted
                upload.analysis_status = "completed"
                await db.commit()
            if upload.extracted_text:
                uploaded_docs.append(f"Document Name: {upload.original_name}\nContent:\n{upload.extracted_text}")

        # Update Project Status to running
        project.status = "running"
        project.progress = 2.0
        project.current_agent = "manager"
        await db.commit()

        user_goal = project.goal
        project_title = project.title
        category = project.category

    # 2. Build AgentForgeState dictionary
    state: AgentForgeState = {
        "project_id": project_id,
        "user_id": user_id,
        "user_goal": user_goal,
        "project_title": project_title,
        "category": category,
        "uploaded_documents": uploaded_docs,
        
        "execution_plan": None,
        "subtasks": None,
        "market_research": None,
        "industry_trends": None,
        "customer_segments": None,
        "swot_analysis": None,
        "market_size": None,
        
        "competitors": None,
        "competitive_landscape": None,
        "competitive_advantages": None,
        
        "business_model_canvas": None,
        "revenue_model": None,
        "pricing_strategy": None,
        "growth_strategy": None,
        "product_roadmap": None,
        
        "budget_estimation": None,
        "revenue_forecast": None,
        "roi_analysis": None,
        "break_even_analysis": None,
        "financial_risks": None,
        
        "marketing_strategy": None,
        "seo_plan": None,
        "social_media_campaigns": None,
        "content_calendar": None,
        "brand_positioning": None,
        
        "risk_register": None,
        "risk_score": None,
        "mitigation_plan": None,
        
        "quality_score": 0.0,
        "qa_feedback": [],
        "approved": False,
        
        "executive_summary": None,
        "final_report": None,
        "report_id": None,
        
        "current_agent": "manager",
        "completed_agents": [],
        "errors": [],
        "start_time": datetime.utcnow().isoformat(),
        "messages": [],
    }

    # 3. Create and compile state graph
    workflow = StateGraph(AgentForgeState)
    workflow.add_node("manager", manager_node)
    workflow.add_node("research", research_node)
    workflow.add_node("competitor", competitor_node)
    workflow.add_node("strategy", strategy_node)
    workflow.add_node("financial", financial_node)
    workflow.add_node("marketing", marketing_node)
    workflow.add_node("risk", risk_node)
    workflow.add_node("qa", qa_node)
    workflow.add_node("report", report_node)

    workflow.set_entry_point("manager")
    workflow.add_edge("manager", "research")
    workflow.add_edge("research", "competitor")
    workflow.add_edge("competitor", "strategy")
    workflow.add_edge("strategy", "financial")
    workflow.add_edge("financial", "marketing")
    workflow.add_edge("marketing", "risk")
    workflow.add_edge("risk", "qa")
    workflow.add_edge("qa", "report")
    workflow.add_edge("report", END)

    app_graph = workflow.compile()

    # 4. Execute workflow
    try:
        await app_graph.ainvoke(state)
        
        # Success actions: deduct credits, update project, add notification
        async with AsyncSessionLocal() as db:
            proj_result = await db.execute(select(Project).where(Project.id == uuid.UUID(project_id)))
            project = proj_result.scalar_one_or_none()
            if project:
                project.status = "completed"
                project.progress = 100.0
                project.completed_at = datetime.utcnow()
                project.credits_used = 10
                
            user_result = await db.execute(select(User).where(User.id == uuid.UUID(user_id)))
            user = user_result.scalar_one_or_none()
            if user:
                user.ai_credits = max(0, user.ai_credits - 10)
            
            # Create user notification
            notification = Notification(
                user_id=uuid.UUID(user_id),
                title="Project Complete! 📊",
                message=f"Agent workflow for '{project_title}' finished successfully. Your Intelligence Report is ready.",
                type="success",
                action_url=f"/reports"
            )
            db.add(notification)
            
            activity = Activity(
                user_id=uuid.UUID(user_id),
                project_id=uuid.UUID(project_id),
                action="workflow_completed",
                description=f"Multi-agent workflow for '{project_title}' completed.",
            )
            db.add(activity)
            
            await db.commit()

        # Final broadcast
        await broadcast_ws(project_id, {
            "type": "workflow_complete",
            "project_id": project_id,
            "message": "All agents have completed execution."
        })
        logger.info(f"Workflow completed successfully for project {project_id}.")

    except Exception as e:
        logger.error(f"Workflow failed for project {project_id}: {e}")
        async with AsyncSessionLocal() as db:
            proj_result = await db.execute(select(Project).where(Project.id == uuid.UUID(project_id)))
            project = proj_result.scalar_one_or_none()
            if project:
                project.status = "failed"
                await db.commit()

            # Create failure notification
            notification = Notification(
                user_id=uuid.UUID(user_id),
                title="Project Execution Failed ⚠️",
                message=f"Agent workflow for '{project_title}' failed. Details: {str(e)}",
                type="error"
            )
            db.add(notification)
            await db.commit()

        await broadcast_ws(project_id, {
            "type": "error",
            "project_id": project_id,
            "message": f"Workflow execution encountered an error: {str(e)}"
        })
